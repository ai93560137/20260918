#!/usr/bin/env python3
"""港股中文名：港交所官方「證券名單」（中文版 Excel）→ data/equities/hk/_hkex_listed.json。

    python3 scripts/fetch_hkex_names.py                 # 在 GitHub Actions 上跑（本機網路白名單擋港交所）
    python3 scripts/fetch_hkex_names.py --file ListOfSecurities_c.xlsx   # 本機解析已下載的檔

輸出 {"asof": 名單更新日, "listed": {"0700.HK": {"name_zh": "騰訊控股", "name_en": ..., "category": ...}}}，
給每日篩選（scripts/daily_topdown.py）顯示中文名；英文名同時從英文版名單取（抓不到就留空）。
"""
import argparse
import io
import json
import re
import sys
import urllib.request
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "data" / "equities" / "hk" / "_hkex_listed.json"
URL_ZH = "https://www.hkex.com.hk/chi/services/trading/securities/securitieslists/ListOfSecurities_c.xlsx"
URL_EN = "https://www.hkex.com.hk/eng/services/trading/securities/securitieslists/ListOfSecurities.xlsx"
UA = "Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124.0 Safari/537.36"


def download(url: str) -> bytes:
    req = urllib.request.Request(url, headers={"User-Agent": UA})
    with urllib.request.urlopen(req, timeout=60) as r:
        return r.read()


def parse(xlsx: bytes, code_keys: tuple[str, ...], name_keys: tuple[str, ...]) -> tuple[str, dict[str, dict]]:
    """回傳 (更新日字串, {代碼: {name, category}})。表頭列用關鍵字找（檔案前幾列是標題/更新日）。"""
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(xlsx), read_only=True, data_only=True)
    ws = wb.worksheets[0]
    asof, header, out = "", None, {}
    for row in ws.iter_rows(values_only=True):
        cells = ["" if v is None else str(v).strip() for v in row]
        if header is None:
            joined = " ".join(cells)
            m = re.search(r"(\d{1,2}/\d{1,2}/\d{4})", joined)
            if m and not asof:
                asof = m.group(1)
            ci = next((i for i, c in enumerate(cells) if any(k in c for k in code_keys)), None)
            ni = next((i for i, c in enumerate(cells) if any(k in c for k in name_keys)), None)
            if ci is not None and ni is not None:
                cat = next((i for i, c in enumerate(cells) if c in ("分類", "Category")), None)
                header = (ci, ni, cat)
            continue
        ci, ni, cat = header
        code = cells[ci] if ci < len(cells) else ""
        if not code.isdigit():
            continue
        out[f"{int(code):04d}.HK" if int(code) < 10000 else f"{int(code)}.HK"] = {
            "name": cells[ni] if ni < len(cells) else "",
            "category": cells[cat] if cat is not None and cat < len(cells) else ""}
    if header is None:
        raise SystemExit("找不到表頭（代號/名稱欄），港交所可能改了格式；前幾列：" + repr(
            [r for _, r in zip(range(6), ws.iter_rows(values_only=True))]))
    return asof, out


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--file", type=Path, help="本機的中文版 xlsx（不下載）")
    args = ap.parse_args()
    zh_bytes = args.file.read_bytes() if args.file else download(URL_ZH)
    asof, zh = parse(zh_bytes, ("股份代號", "代號"), ("股份名稱", "名稱"))
    en: dict[str, dict] = {}
    if not args.file:
        try:
            _, en = parse(download(URL_EN), ("Stock Code",), ("Name of Securities",))
        except Exception as exc:   # 英文名只是附帶，失敗不影響中文名
            print(f"WARN 英文版名單抓取失敗：{exc}", file=sys.stderr)
    if len(zh) < 1000:
        raise SystemExit(f"只解析到 {len(zh)} 檔，太少，不覆寫舊檔")
    listed = {t: {"name_zh": v["name"], "name_en": en.get(t, {}).get("name", ""), "category": v["category"]}
              for t, v in sorted(zh.items())}
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps({"asof": asof, "source": URL_ZH, "listed": listed},
                              ensure_ascii=False, indent=1, sort_keys=True) + "\n", encoding="utf-8")
    sample = {t: listed[t]["name_zh"] for t in ("0005.HK", "0700.HK", "2359.HK") if t in listed}
    print(f"港交所證券名單（{asof}）{len(listed)} 檔 -> {OUT}；例：{sample}")


if __name__ == "__main__":
    main()
